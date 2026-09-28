"""Contracts for the bilingual Skyward data-notes page."""
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "app" / "templates"
STATIC = ROOT / "app" / "static"


def test_about_and_api_titles_reuse_home_page_title_class():
    index = (TEMPLATES / "index.html").read_text(encoding="utf-8")
    about = (TEMPLATES / "about.html").read_text(encoding="utf-8")
    api = (TEMPLATES / "api.html").read_text(encoding="utf-8")
    assert 'class="page-title" data-i18n="homeTitle"' in index
    assert 'class="page-title" data-i18n="aboutTitle"' in about
    assert 'class="page-title" data-i18n="apiTitle"' in api


def test_data_notes_have_usage_and_separate_current_goal_future_scope():
    about = (TEMPLATES / "about.html").read_text(encoding="utf-8")
    script = (STATIC / "about.js").read_text(encoding="utf-8")
    for key in (
        "currentStageLabel",
        "goalStageLabel",
        "futureStageLabel",
        "usageTitle",
        "usageStep1Body",
        "usageStep2Body",
        "usageStep3Body",
        "usageStep4Body",
    ):
        assert f'data-about-i18n="{key}"' in about
        assert key in script
    assert "只有确认后才加载所选源表数据并绘制标记" in script
    assert "geometric-feasibility candidates" in script
    assert "不构成观测批准" in script
    assert "天气与大气" in script and "机械轨迹限位" in script and "联合观测证据" in script
    assert 'data-about-i18n="observingPlanTitle"' in about
    assert 'data-about-i18n="observingPlanBody"' in about
    assert 'data-about-i18n="observingPlanFormat"' in about
    assert "完整源窗口区域只保留一组计划操作" in script
    assert "one shared set of plan actions" in script
    assert "观测计划仅在浏览器本地生成" in script
    assert "The observing plan is generated only in the browser" in script
    api = (TEMPLATES / "api.html").read_text(encoding="utf-8")
    app_script = (STATIC / "app.js").read_text(encoding="utf-8")
    assert 'data-i18n="apiPlanWindowNote"' in api
    assert "多窗口计划逐窗口调用" in app_script
    assert "multi-window plans call once per window" in app_script


def test_catalogue_guide_is_one_select_with_five_targeted_choices():
    about = (TEMPLATES / "about.html").read_text(encoding="utf-8")
    select = re.search(r'<select id="about-catalogue-select"[^>]*>(.*?)</select>', about, re.S)
    assert select is not None
    assert " multiple" not in select.group(0)
    values = re.findall(r'<option value="([^"]+)">', select.group(1))
    assert values == ["1lhaaso", "fermi-fl16y", "fermi-3fhl", "tevcat", "gaia-dr3"]
    assert 'data-i18n="catalogueTitle"' not in about
    assert 'data-about-i18n="catalogueGuideTitle"' in about


def test_data_notes_describe_xlsx_upload_contracts_and_examples():
    about = (TEMPLATES / "about.html").read_text(encoding="utf-8")
    about_js = (STATIC / "about.js").read_text(encoding="utf-8")
    for key in (
        "csvUploadTitle", "csvUploadIntro", "csvUploadName", "sourceCsvTitle", "targetListCsvTitle", "targetListCsvDownload", "targetListCsvBody", "targetListCsvExample", "sourceCsvDownload",
        "sourceCsvBody", "alternativeCsvTitle", "alternativeCsvDownload", "alternativeCsvBody", "csvUploadLimits", "csvUploadAliases",
    ):
        assert f'data-about-i18n="{key}"' in about
        assert key in about_js
    assert 'download="user-source-catalogue.xlsx"' in about
    assert 'download="target-source-list.xlsx"' in about
    assert 'download="alternative-source-catalogue.xlsx"' in about
    expected = {
        "user-source-catalogue.xlsx": ("name", "Crab"),
        "target-source-list.xlsx": ("name", "Crab target"),
        "alternative-source-catalogue.xlsx": ("name", "Alternative A"),
    }
    for filename, (header, first_value) in expected.items():
        workbook = load_workbook(STATIC / "examples" / filename, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
        assert rows[0][0] == header
        assert rows[1][0] == first_value
        assert len(rows) >= 3
    assert "1,000 rows and 512 KiB" in about_js
    assert "8 MiB" in about_js
    assert "Observing plans" in about
    assert "solid black" not in about
    assert "纯黑色" not in about
    assert "最多 1,000 行、512 KiB" in about_js


    script = (STATIC / "about.js").read_text(encoding="utf-8")
    assert "Installed row count is provided by the catalogue API" in script
    assert "363-row snapshot · cutoff 2026-09-25" in script
    assert "500 hard maximum per request" in script
    assert "Installed row count is provided by the catalogue API" in script
    assert "Localization uncertainty is not physical extension" in script
    assert "223 upstream Extended: No rows are zero-radius point sources" in script
    assert "140 Extended: Yes rows" in script
    assert "nearest N sources by angular distance" in script
    assert "does not appear in an all-sky catalogue picker or layer" in script
    assert "本地记录数由源表接口提供" in script
    assert "本地快照 363 条 · 截止 2026-09-25" in script
    assert "2026-09-25（北京时间，UTC+8）" in script
    assert "2026-09-25 Beijing time (UTC+8)" in script
    assert "20260924 按北京时间（UTC+8）日历日期定义" in script
    assert "硬上限 500 条/次" in script
    assert "仅在计算结果页的局部视场图中" in script


def test_about_assets_are_local_and_follow_shared_language_event():
    base = (TEMPLATES / "base.html").read_text(encoding="utf-8")
    script = (STATIC / "about.js").read_text(encoding="utf-8")
    stylesheet = (STATIC / "about.css").read_text(encoding="utf-8")
    assert "path='/about.css'" in base
    assert "path='/about.js'" in base
    assert 'localStorage.getItem("skyward.language")' in script
    assert 'document.addEventListener("skyward:language-change", applyLanguage)' in script
    assert "fetch(\"/api/v1/catalogues\")" in script
    assert "@media (max-width: 760px)" in stylesheet
    about = (TEMPLATES / "about.html").read_text(encoding="utf-8")
    app_script = (STATIC / "app.js").read_text(encoding="utf-8")
    assert 'class="legend-cross"' in about
    assert "加粗十字符号表示按需查询的 Gaia 候选星" in app_script
    assert "A thicker cross marks an on-demand Gaia candidate" in app_script
